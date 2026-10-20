"""
Motor de reportes binarios del módulo Personal (SRP/LSP).

Concentra la maquetación de exportación en Excel y PDF que antes vivía
inline en ``apps.personal.views``:

* ``generar_excel_personal`` / ``generar_pdf_personal``: listado de personal
  (reciben el queryset ya filtrado por la vista).
* ``generar_excel_alumnos_representantes`` / ``generar_pdf_alumnos_representantes``:
  listado de alumnos y sus representantes legales.

Todas devuelven estrictamente ``io.BytesIO`` con el cursor en posición 0
(buffer en memoria, sin I/O de disco): el transporte HTTP (FileResponse,
headers, content-type) es exclusivamente responsabilidad de las vistas —
así cualquier vista puede consumir estos generadores de forma
intercambiable (Principio de Sustitución de Liskov).
"""

from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Cabeceras compartidas de las hojas de Excel (idénticas en ambos listados).
_FUENTE_CABECERA = Font(name="Arial", size=11, bold=True, color="FFFFFF")
_RELLENO_CABECERA = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")


def _ajustar_ancho_columnas(ws) -> None:
    """Ajusta el ancho de cada columna al contenido máximo (mín. 12)."""
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def generar_excel_personal(queryset) -> BytesIO:
    """Listado de personal en Excel (.xlsx) a partir del queryset recibido."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Personal"

    encabezados = ['Cédula', 'Primer Nombre', 'Primer Apellido', 'Categoría', 'Teléfono', 'Correo']
    ws.append(encabezados)

    for col_num in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col_num)
        celda.font = _FUENTE_CABECERA
        celda.fill = _RELLENO_CABECERA

    for p in queryset:
        telefono = getattr(p, 'telefono', '') or (getattr(p.id_usuario, 'telefono', '') if p.id_usuario else '')
        correo = getattr(p.id_usuario, 'email', '') or ''
        ws.append([
            p.cedula,
            p.primer_nombre,
            p.primer_apellido,
            p.categoria,
            telefono or 'No registrado',
            correo or 'No registrado',
        ])

    _ajustar_ancho_columnas(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_pdf_personal(queryset, categoria=None) -> BytesIO:
    """
    Listado de personal en PDF a partir del queryset recibido.

    ``categoria`` (opcional) segmenta el título: "LISTADO DE PERSONAL — <CATEGORIA>".
    """
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []
    styles = getSampleStyleSheet()

    estilo_titulo = ParagraphStyle(
        'TituloPersonal',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1F4E78'),
        alignment=1
    )

    # Título segmentado por la categoría solicitada (si viene)
    titulo = "LISTADO DE PERSONAL"
    if categoria:
        titulo = f"LISTADO DE PERSONAL — {categoria.upper()}"

    story.append(Paragraph(titulo, estilo_titulo))
    story.append(Spacer(1, 16))

    datos_tabla = [['Cédula', 'Nombre', 'Categoría', 'Teléfono', 'Correo']]
    for p in queryset:
        telefono = getattr(p, 'telefono', '') or (getattr(p.id_usuario, 'telefono', '') if p.id_usuario else '')
        correo = getattr(p.id_usuario, 'email', '') or ''
        datos_tabla.append([
            p.cedula,
            f"{p.primer_nombre} {p.primer_apellido}",
            p.categoria,
            telefono or 'No registrado',
            correo or 'No registrado',
        ])

    tabla_pdf = Table(datos_tabla, colWidths=[90, 150, 140, 110, 110])
    tabla_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F2F4F7')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
    ]))

    story.append(tabla_pdf)
    doc.build(story)
    buffer.seek(0)
    return buffer


def generar_excel_alumnos_representantes(relaciones) -> BytesIO:
    """Parejas alumno–representante en Excel (.xlsx) a partir del queryset recibido."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alumnos - Representantes"

    encabezados = ['Alumno', 'Representante', 'Teléfono', 'Parentesco', 'Es Principal']
    ws.append(encabezados)

    for col_num in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col_num)
        celda.font = _FUENTE_CABECERA
        celda.fill = _RELLENO_CABECERA

    for relacion in relaciones:
        ws.append([
            f"{relacion.alumno.primer_nombre} {relacion.alumno.primer_apellido}",
            f"{relacion.representante.primer_nombre} {relacion.representante.primer_apellido}",
            relacion.representante.telefono_principal or 'No registrado',
            relacion.parentesco or 'No especificado',
            "Sí" if relacion.es_principal else "No",
        ])

    _ajustar_ancho_columnas(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_pdf_alumnos_representantes(relaciones) -> BytesIO:
    """Parejas alumno–representante en PDF a partir del queryset recibido."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []

    styles = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        'TituloReporte',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1F4E78'),
        alignment=1
    )

    story.append(Paragraph("PLATAFORMA ESCOLAR - CONTROL DE INSCRIPCIONES", estilo_titulo))
    story.append(Spacer(1, 20))

    datos_tabla = [['Alumno', 'Representante', 'Teléfono', 'Parentesco', 'Es Principal']]

    for relacion in relaciones:
        datos_tabla.append([
            f"{relacion.alumno.primer_nombre} {relacion.alumno.primer_apellido}",
            f"{relacion.representante.primer_nombre} {relacion.representante.primer_apellido}",
            relacion.representante.telefono_principal or 'No registrado',
            relacion.parentesco or 'No especificado',
            "Principal" if relacion.es_principal else "Secundario",
        ])

    tabla_pdf = Table(datos_tabla, colWidths=[150, 150, 120, 100, 100])
    tabla_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 11),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F2F4F7')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
    ]))

    story.append(tabla_pdf)
    doc.build(story)
    buffer.seek(0)
    return buffer
