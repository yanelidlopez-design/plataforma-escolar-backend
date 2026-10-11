from rest_framework import viewsets, permissions
from rest_framework.decorators import api_view, permission_classes, action
from drf_yasg.utils import swagger_auto_schema
from drf_yasg import openapi
from django.http import HttpResponse
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from .models import Profesores, RepresentantesLegales, AlumnoRepresentante
from .serializers import ProfesoresSerializer, RepresentantesLegalesSerializer, AlumnoRepresentanteSerializer


@swagger_auto_schema(tags=['Personal'])
class PersonalViewSet(viewsets.ModelViewSet):
    queryset = Profesores.objects.select_related('id_usuario').all()
    serializer_class = ProfesoresSerializer
    search_fields = ['cedula', 'primer_nombre', 'segundo_nombre', 'primer_apellido', 'segundo_apellido', 'categoria']

    def get_queryset(self):
        """
        Filtrado manual por query params (django-filter NO está disponible):
        GET /api/personal/personal/?categoria=Profesor
        Categorías inválidas se ignoran limpiamente (fail-soft en lectura);
        la validación estricta de categorías ocurre en el POST (serializer).
        """
        qs = super().get_queryset()
        categoria = (self.request.query_params.get('categoria') or '').strip()
        if categoria:
            validas = {codigo for codigo, _ in Profesores.CATEGORIA_CHOICES}
            if categoria in validas:
                qs = qs.filter(categoria=categoria)
        return qs

    def perform_create(self, serializer):
        """
        La tabla legacy `profesores` exige id_usuario (NOT NULL, UNIQUE, FK a usuarios).
        Si el payload no trae id_usuario, se crea automáticamente el Usuario vinculado
        (con rol derivado de la categoría) de forma atómica para no dejar registros huérfanos.
        """
        data = dict(serializer.validated_data)
        id_usuario = data.pop('id_usuario', None)

        if id_usuario:
            serializer.save(id_usuario=id_usuario)
            return

        from django.db import transaction
        from django.contrib.auth.hashers import make_password
        from django.utils.crypto import get_random_string
        from apps.core.models import Usuarios

        cedula = data.get('cedula') or ''
        primer_nombre = data.get('primer_nombre') or ''
        primer_apellido = data.get('primer_apellido') or ''
        categoria = data.get('categoria') or 'Profesor'

        # Rol en usuarios según la categoría del personal
        categoria_lower = categoria.lower()
        if 'director' in categoria_lower or 'coordinador' in categoria_lower:
            rol = 'directivo'
        elif categoria_lower == 'administrativo':
            rol = 'secretaria'
        else:
            rol = 'profesor'

        # Email único (la columna usuarios.email es UNIQUE en la BD legacy)
        email = (self.request.data.get('email') or '').strip()
        if not email or Usuarios.objects.filter(email=email).exists():
            email = self._generar_email_unico(cedula)

        with transaction.atomic():
            usuario = Usuarios.objects.create(
                nombre_usuario=self._generar_nombre_usuario(primer_nombre, primer_apellido, cedula),
                clave_hash=make_password(get_random_string(24)),
                email=email,
                rol=rol,
                activo=True,
            )
            serializer.save(id_usuario=usuario)

    @staticmethod
    def _generar_email_unico(cedula):
        from apps.core.models import Usuarios
        base = f"{(cedula or 'personal').replace('-', '').lower()}@plantel.local"
        local, dominio = base.split('@')
        email = base
        contador = 1
        while Usuarios.objects.filter(email=email).exists():
            email = f"{local}{contador}@{dominio}"
            contador += 1
        return email

    @staticmethod
    def _generar_nombre_usuario(primer_nombre, primer_apellido, cedula):
        base = (f"{primer_nombre.split()[0] if primer_nombre else ''}"
                f".{primer_apellido.split()[0] if primer_apellido else ''}").strip('.').lower().replace(' ', '')
        if not base:
            base = f"user{(cedula or '0').replace('-', '')}"
        from apps.core.models import Usuarios
        nombre = base
        contador = 1
        while Usuarios.objects.filter(nombre_usuario=nombre).exists():
            nombre = f"{base}{contador}"
            contador += 1
        return nombre

    @swagger_auto_schema(
        method='GET',
        operation_description='Exportar personal filtrado por categoría en formato Excel (.xlsx)',
        manual_parameters=[
            openapi.Parameter('categoria', openapi.IN_QUERY, description='Categoría del personal (opcional)', type=openapi.TYPE_STRING),
        ],
        responses={200: openapi.Response('Archivo Excel')}
    )
    @action(detail=False, methods=['get'], url_path='exportar-excel')
    def exportar_excel(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Personal"

        fuente_cabecera = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        relleno_azul = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

        encabezados = ['Cédula', 'Primer Nombre', 'Primer Apellido', 'Categoría', 'Teléfono', 'Correo']
        ws.append(encabezados)

        for col_num in range(1, len(encabezados) + 1):
            celda = ws.cell(row=1, column=col_num)
            celda.font = fuente_cabecera
            celda.fill = relleno_azul

        # get_queryset aplica el filtro ?categoria= recibido en la request
        personal = self.get_queryset()

        for p in personal:
            telefono = getattr(p, 'telefono', '') or (getattr(p.id_usuario, 'telefono', '') if p.id_usuario else '')
            correo = getattr(p.id_usuario, 'email', '') or ''

            ws.append([
                p.cedula,
                p.primer_nombre,
                p.primer_apellido,
                p.categoria,
                telefono or 'No registrado',
                correo or 'No registrado',
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="listado_personal.xlsx"'

        wb.save(response)
        return response

    @swagger_auto_schema(
        method='GET',
        operation_description='Exportar personal filtrado por categoría en formato PDF',
        manual_parameters=[
            openapi.Parameter('categoria', openapi.IN_QUERY, description='Categoría del personal (opcional)', type=openapi.TYPE_STRING),
        ],
        responses={200: openapi.Response('Archivo PDF')}
    )
    @action(detail=False, methods=['get'], url_path='exportar-pdf')
    def exportar_pdf(self, request):
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="listado_personal.pdf"'

        doc = SimpleDocTemplate(response, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        story = []
        styles = getSampleStyleSheet()

        estilo_titulo = ParagraphStyle(
            'TituloPersonal',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1F4E78'),
            alignment=1
        )

        # Título segmentado por la categoría solicitada (si viene)
        categoria = (request.query_params.get('categoria') or '').strip()
        titulo = "LISTADO DE PERSONAL"
        if categoria:
            titulo = f"LISTADO DE PERSONAL — {categoria.upper()}"

        story.append(Paragraph(titulo, estilo_titulo))
        story.append(Spacer(1, 16))

        # get_queryset aplica el filtro ?categoria= recibido en la request
        datos_tabla = [['Cédula', 'Nombre', 'Categoría', 'Teléfono', 'Correo']]
        for p in self.get_queryset():
            telefono = getattr(p, 'telefono', '') or (getattr(p.id_usuario, 'telefono', '') if p.id_usuario else '')
            correo = getattr(p.id_usuario, 'email', '') or ''
            datos_tabla.append([
                p.cedula,
                f"{p.primer_nombre} {p.primer_apellido}",
                p.categoria,
                telefono or 'No registrado',
                correo or 'No registrado',
            ])

        tabla_pdf = Table(datos_tabla, colWidths=[90, 150, 140, 110, 110])
        tabla_pdf.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F2F4F7')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ]))

        story.append(tabla_pdf)
        doc.build(story)
        return response

    @action(detail=False, methods=['get'], url_path='exportar-csv')
    def exportar_csv(self, request):
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="listado_personal.csv"'

        writer = csv.writer(response)
        writer.writerow(['Cedula', 'Primer Nombre', 'Segundo Nombre', 'Primer Apellido', 'Segundo Apellido', 'Categoria', 'Correo'])

        # get_queryset aplica el filtro ?categoria= recibido en la request
        for p in self.get_queryset():
            correo = getattr(p.id_usuario, 'email', '') or ''
            writer.writerow([
                p.cedula,
                p.primer_nombre,
                p.segundo_nombre or '',
                p.primer_apellido,
                p.segundo_apellido or '',
                p.categoria,
                correo or 'No registrado',
            ])

        return response


@swagger_auto_schema(
    method='GET',
    operation_description='Exportar alumnos y representantes en formato CSV',
    tags=['Exportaciones'],
    responses={200: openapi.Response('Archivo CSV')}
)
@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def exportar_alumnos_representantes_csv(request):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="alumnos_y_representantes.csv"'

    writer = csv.writer(response)
    writer.writerow(['Alumno', 'Representante', 'Teléfono', 'Parentesco', 'Es Principal'])

    relaciones = AlumnoRepresentante.objects.select_related('alumno', 'representante')
    for relacion in relaciones:
        writer.writerow([
            f"{relacion.alumno.primer_nombre} {relacion.alumno.primer_apellido}",
            f"{relacion.representante.primer_nombre} {relacion.representante.primer_apellido}",
            relacion.representante.telefono_principal or 'No registrado',
            relacion.parentesco or 'No especificado',
            "Sí" if relacion.es_principal else "No"
        ])

    return response


@swagger_auto_schema(
    method='GET',
    operation_description='Exportar alumnos y representantes en formato Excel (.xlsx)',
    tags=['Exportaciones'],
    responses={200: openapi.Response('Archivo Excel')}
)
@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def exportar_alumnos_representantes_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alumnos - Representantes"

    fuente_cabecera = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    relleno_azul = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    encabezados = ['Alumno', 'Representante', 'Teléfono', 'Parentesco', 'Es Principal']
    ws.append(encabezados)

    for col_num in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col_num)
        celda.font = fuente_cabecera
        celda.fill = relleno_azul

    relaciones = AlumnoRepresentante.objects.select_related('alumno', 'representante')

    for relacion in relaciones:
        ws.append([
            f"{relacion.alumno.primer_nombre} {relacion.alumno.primer_apellido}",
            f"{relacion.representante.primer_nombre} {relacion.representante.primer_apellido}",
            relacion.representante.telefono_principal or 'No registrado',
            relacion.parentesco or 'No especificado',
            "Sí" if relacion.es_principal else "No"
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="alumnos_y_representantes.xlsx"'

    wb.save(response)
    return response


@swagger_auto_schema(
    method='GET',
    operation_description='Exportar alumnos y representantes en formato PDF',
    tags=['Exportaciones'],
    responses={200: openapi.Response('Archivo PDF')}
)
@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def exportar_alumnos_representantes_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="alumnos_y_representantes.pdf"'

    doc = SimpleDocTemplate(response, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []

    styles = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        'TituloReporte',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1F4E78'),
        alignment=1
    )

    story.append(Paragraph("PLATAFORMA ESCOLAR - CONTROL DE INSCRIPCIONES", estilo_titulo))
    story.append(Spacer(1, 20))

    datos_tabla = [['Alumno', 'Representante', 'Teléfono', 'Parentesco', 'Es Principal']]

    relaciones = AlumnoRepresentante.objects.select_related('alumno', 'representante')
    for relacion in relaciones:
        datos_tabla.append([
            f"{relacion.alumno.primer_nombre} {relacion.alumno.primer_apellido}",
            f"{relacion.representante.primer_nombre} {relacion.representante.primer_apellido}",
            relacion.representante.telefono_principal or 'No registrado',
            relacion.parentesco or 'No especificado',
            "Principal" if relacion.es_principal else "Secundario"
        ])

    tabla_pdf = Table(datos_tabla, colWidths=[150, 150, 120, 100, 100])
    tabla_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 11),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F2F4F7')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')])
    ]))

    story.append(tabla_pdf)

    doc.build(story)
    return response


@swagger_auto_schema(tags=['Representantes Legales'])
class RepresentantesLegales(viewsets.ModelViewSet):
    queryset = RepresentantesLegales.objects.select_related('id_direccion_residencia').all()
    serializer_class = RepresentantesLegalesSerializer
    search_fields = ['cedula', 'primer_nombre', 'segundo_nombre', 'primer_apellido', 'segundo_apellido']


@swagger_auto_schema(tags=['Alumno-Representante'])
class AlumnoRepresentanteViewSet(viewsets.ModelViewSet):
    queryset = AlumnoRepresentante.objects.all()
    serializer_class = AlumnoRepresentanteSerializer
    filterset_fields = ['id_alumno', 'id_representante', 'es_principal']
